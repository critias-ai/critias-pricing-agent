"""Persistent SQL session for the analyst: one DuckDB connection with typed views over the prepared sources.
JSON lines in ({"id","sql"}), JSON lines out. Facts only; this module has no financial definitions."""
import argparse
import csv
import io
import itertools
import json
import math
import os
import re
import sys
import threading
import time
from decimal import Decimal
from pathlib import Path

import duckdb

sys.path.insert(0, str(Path(__file__).resolve().parent))
from data_workspace import ImportOptions, attach_source, connect_analysis, create_typed_view, header_row, prepare_source, profile_prepared, text_encoding
from same_rows import mismatches, note, stop_message
from peer_ratio import PeerRatio
from offsets import Offsets

# Statement kinds the analyst may run. Everything else (ATTACH, SET, PRAGMA, INSTALL, LOAD, CALL, ...) is refused.
ALLOWED = re.compile(
    r'^\s*(select|with|from|describe|show|summarize|explain|values|'
    r'create\s+(or\s+replace\s+)?(temp(orary)?\s+)?(view|table)|drop\s+(view|table)(\s+if\s+exists)?|copy)\b',
    re.I,
)
MAX_ROWS = 100


def alias_for(path):
    stem = re.sub(r'[^a-z0-9]+', '_', Path(path).stem.lower()).strip('_')
    return stem if stem and stem[0].isalpha() else 't_' + stem


def shown(value):
    """Decimals without the storage padding: at least two places, no trailing zeros beyond that."""
    if isinstance(value, float) and not math.isfinite(value):
        return None  # NaN and infinity are not JSON; they show as NULL
    if not isinstance(value, Decimal):
        return value
    text = f'{value:f}'
    if '.' not in text:
        return text
    whole, places = text.split('.')
    return whole + '.' + (places.rstrip('0') + '00')[: max(2, len(places.rstrip('0')))]


def definitions(conn):
    rows = conn.execute('SELECT view_name, sql FROM duckdb_views() WHERE NOT internal').fetchall()
    rows += conn.execute('SELECT table_name, sql FROM duckdb_tables() WHERE NOT internal').fetchall()
    return {name: sql for name, sql in rows if not name.startswith('_')}  # _result, _peer_*: the harness's own


# A decimal or a timestamp is never a column name; the same rule as data_findings.DATA_CELL.
DATA_CELL = re.compile(r'^[+-]?[0-9]*[.][0-9]+$|^[0-9]{4}-[0-9]{2}-[0-9]{2}[ T][0-9]{2}:[0-9]{2}')
NO_HEADER = ('{}: no header row (the first row holds data), so its columns are named column1 to column{}. Ask the user '
             'for the column names, or to confirm names found in another file, then define a view that names them.')


def targets(path):
    """(view name, import options) for each table in a source: a CSV, or every sheet of a workbook."""
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True)
        try:
            return [(alias_for(path) + '_' + alias_for(sheet + '.x').removeprefix('t_'), ImportOptions(format='xlsx', sheet=sheet))
                    for sheet in book.sheetnames]
        finally:
            book.close()
    with open(path, encoding=text_encoding(path), newline='') as f:
        sample = f.read(65536)
    delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
    row = header_row(sample, delimiter)
    first = next(itertools.islice(csv.reader(io.StringIO(sample), delimiter=delimiter), row - 1, None), [])
    headerless = any(DATA_CELL.match(v.strip()) for v in first)
    return [(alias_for(path), ImportOptions(format='csv', header_row=0 if headerless else row))]


def attach_views(conn, root, cache, sources):
    """Attach each CSV and workbook sheet read-only and give it a typed view named after the file (and sheet)."""
    tables, notes = [], []
    for rel in sources:
        try:
            found = targets(Path(root) / rel)
        except Exception as error:
            notes.append(f'{rel}: not available ({str(error)[:160]})')
            continue
        for name, options in found:
            try:
                prepared = prepare_source(Path(root) / rel, Path(cache), options)
                attach_source(conn, prepared, 'src_' + name)
                try:
                    create_typed_view(conn, 'src_' + name, profile_prepared(prepared), view_name=name)
                    typed = True
                except ValueError as error:
                    conn.execute(f'CREATE OR REPLACE VIEW "{name}" AS SELECT * FROM "src_{name}".raw')
                    notes.append(f'{name}: kept as text ({str(error)[:160]})')
                    typed = False
                columns = [(r[0], r[1]) for r in conn.execute(f'DESCRIBE "{name}"').fetchall()]
                if not options.header_row:
                    notes.append(NO_HEADER.format(name, len(columns)))
                if options.header_row > 1:
                    notes.append(f'{name}: the header is on row {options.header_row}; the rows above it (a title or date range) were left out.')
                tables.append({'name': name, 'rows': prepared.row_count, 'typed': typed, 'file': rel,
                               **({'sheet': options.sheet} if options.sheet else {}),
                               'columns': [f'{c} {t}' for c, t in columns]})
            except Exception as error:  # an unreadable source must not take the others down
                notes.append(f'{rel}{"#" + options.sheet if options.sheet else ""}: not available ({str(error)[:160]})')
    return tables, notes


def open_session(stage, memory, threads):
    """The conversation's database. It is kept between turns, so the analyst's views and tables persist; the workflow
    deletes the whole stage folder when a conversation starts afresh or its sources change. The previous turn's
    session may still be releasing the file, so a lock is retried briefly."""
    for attempt in range(20):
        try:
            return connect_analysis(stage / 'qdb', stage / 'qtmp', memory, threads)
        except duckdb.IOException as error:
            if 'lock' not in str(error).lower() or attempt == 19:
                raise
            time.sleep(0.25)


def peer_columns(root, cache, sources):
    """Column and key pairs data findings judged well described by the key's median. Findings are cached by source
    content and were computed before the turn, so this is a cache read; any failure means no peer check."""
    try:
        from data_findings import detect
        here = os.getcwd()
        os.chdir(root)  # findings are keyed by the workspace-relative paths the workflow used
        try:
            return detect(sources, cache)[1].get('peers', [])
        finally:
            os.chdir(here)
    except Exception:
        return []


def setup(args):
    stage = Path(args.stage).resolve()
    os.chdir(stage)  # relative COPY targets land in the stage folder
    conn = open_session(stage, args.memory, args.threads)
    sources = json.loads(args.sources)
    tables, notes = attach_views(conn, args.root, args.cache, sources)
    peers = None
    if not args.no_peer_ratio:
        views = {t['file']: t['name'] for t in tables if t['typed']}
        peers = PeerRatio(conn, peer_columns(args.root, args.cache, sources), views)
    # Sources are attached; from here the SQL can only touch this stage folder.
    conn.execute(f"SET allowed_directories=['{stage.as_posix()}']")
    conn.execute('SET enable_external_access=false')
    conn.execute('SET lock_configuration=true')
    return conn, tables, notes, peers


# A TEMP view or table would end with this session; the analyst's definitions are meant to last the conversation.
TEMPORARY = re.compile(r'^(\s*create\s+(?:or\s+replace\s+)?)temp(?:orary)?\s+(?=(?:view|table)\b)', re.I)
SAVE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}\.csv$')


UNSUMMED = ('Note: this file holds {:,} unsummed source rows. Outlier and returns checks run on sums made in SQL; sums '
            'made in Python are not checked. Group in SQL first unless the script needs single rows.')


COPIED = re.compile(r'^\s*copy\s*\((.*)\)\s*to\s', re.I | re.S)


def unsummed(conn, sql):
    """True when the plan reads source rows (attached src_* databases) and sums or groups nothing."""
    plan = ' '.join(r[1] for r in conn.execute('EXPLAIN ' + sql).fetchall())
    return 'src_' in plan and not re.search(r'GROUP_BY|AGGREGATE', plan)


class SameRows:
    """Session state for the same-rows check: statements already stopped once (sent again unchanged, they run) and
    gaps already noted."""

    def __init__(self):
        self.stopped, self.noted = set(), set()


def run(conn, sql, save=None, same_rows=None, peers=None, offsets=None):
    if len(conn.extract_statements(sql)) != 1:
        raise ValueError('Send exactly one SQL statement per call.')
    if not ALLOWED.match(sql):
        raise ValueError('Only SELECT, WITH, DESCRIBE, SHOW, SUMMARIZE, EXPLAIN, CREATE/DROP VIEW or TABLE and COPY ... TO are allowed.')
    sql = TEMPORARY.sub(r'\1', sql, count=1)
    before = definitions(conn)
    timer = threading.Timer(120, conn.interrupt)
    timer.start()
    started = time.time()
    try:
        found, extra = (mismatches(conn, sql) if same_rows else []), ''
        if any(f['combined'] for f in found):
            statement = ' '.join(sql.split())
            if statement not in same_rows.stopped:
                same_rows.stopped.add(statement)
                return {'columns': [], 'rows': [], 'more': False, 'changed': [], 'seconds': 0,
                        'saved': None, 'savedRows': None, 'stopped': stop_message(found)}
        if same_rows:
            extra = note(found, same_rows.noted)
        if peers:
            extra = ' '.join(filter(None, [extra, peers.note(conn, sql)]))
        if offsets:
            extra = ' '.join(filter(None, [extra, offsets.note(conn, sql)]))
        saved_rows = None
        if save:
            if not SAVE.match(save):
                raise ValueError('save must be a plain file name ending in .csv')
            if not re.match(r'^\s*(select|with|from|values)\b', sql, re.I):
                raise ValueError('save works with a SELECT or WITH query.')
            conn.execute('DROP TABLE IF EXISTS _result')
            conn.execute(f'CREATE TEMP TABLE _result AS {sql}')
            conn.execute(f"COPY _result TO '{save}' (HEADER, DELIMITER ',')")
            saved_rows = conn.execute('SELECT count(*) FROM _result').fetchone()[0]
            if unsummed(conn, sql):
                extra = ' '.join(filter(None, [extra, UNSUMMED.format(saved_rows)]))
            cursor = conn.execute('SELECT * FROM _result')
        else:
            cursor = conn.execute(sql)
            copied = COPIED.match(sql)
            if copied and unsummed(conn, copied.group(1)):
                count = conn.execute(f'SELECT count(*) FROM ({copied.group(1)})').fetchone()[0]
                extra = ' '.join(filter(None, [extra, UNSUMMED.format(count)]))
        columns = [d[0] for d in cursor.description] if cursor.description else []
        rows = cursor.fetchmany(MAX_ROWS + 1) if columns else []
    finally:
        timer.cancel()
    after = definitions(conn)
    changed = [{'name': n, 'change': 'created' if n not in before else 'redefined'} for n, s in after.items()
               if before.get(n) != s]
    changed += [{'name': n, 'change': 'dropped'} for n in before if n not in after]
    return {'columns': columns, 'rows': [[shown(v) for v in r] for r in rows[:MAX_ROWS]], 'more': len(rows) > MAX_ROWS,
            'changed': changed, 'seconds': round(time.time() - started, 2),
            'saved': save if save else None, 'savedRows': saved_rows, 'note': extra or None}


def main():
    parser = argparse.ArgumentParser()
    for name in ('root', 'stage', 'cache', 'sources'):
        parser.add_argument('--' + name, required=True)
    parser.add_argument('--memory', default='4GB')
    parser.add_argument('--threads', type=int, default=4)
    parser.add_argument('--no-same-rows', action='store_true')
    parser.add_argument('--no-peer-ratio', action='store_true')
    parser.add_argument('--no-offsets', action='store_true')
    args = parser.parse_args()
    conn, tables, notes, peers = setup(args)
    same_rows = None if args.no_same_rows else SameRows()
    offsets = None if args.no_offsets else Offsets()
    print(json.dumps({'ready': True, 'tables': tables, 'notes': notes}, default=str), flush=True)
    for line in sys.stdin:
        request = json.loads(line)
        print(json.dumps({'id': request['id'], **answer(conn, request['sql'], request.get('save'), same_rows, peers, offsets)},
                         default=str), flush=True)


def outcome(conn, sql, save, same_rows, peers=None, offsets=None):
    try:
        return {'ok': True, **run(conn, sql, save, same_rows, peers, offsets)}
    except Exception as error:
        return {'ok': False, 'error': f'{type(error).__name__}: {error}'[:2000]}


COMMENTS = re.compile(r'^(?:\s*--[^\n]*(?:\n|$))+')


def answer(conn, sql, save=None, same_rows=None, peers=None, offsets=None):
    """One statement, or several separated by semicolons: those run in order and each keeps its own result or error,
    so independent questions cost the analyst one call instead of one each."""
    try:
        statements = [COMMENTS.sub('', s.query).strip() for s in conn.extract_statements(sql)]
    except Exception as error:
        return {'ok': False, 'error': f'{type(error).__name__}: {error}'[:2000]}
    statements = [s for s in statements if s]
    if len(statements) <= 1:
        return outcome(conn, statements[0] if statements else sql, save, same_rows, peers, offsets)
    if save:
        return {'ok': False, 'error': 'ValueError: save works with one statement per call.'}
    return {'ok': True, 'batch': [dict(outcome(conn, s, None, same_rows, peers, offsets), sql=s) for s in statements]}


if __name__ == '__main__':
    main()
