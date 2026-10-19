"""Print parsing profiles for one CSV/XLSX source as JSON. Facts only; no business meanings."""
import argparse
import json
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
from data_workspace import ImportOptions,prepare_source,profile_prepared

def sheet_names(path):
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True)
    try:return list(book.sheetnames)
    finally:book.close()

def main():
    parser=argparse.ArgumentParser(description='Profile one raw source; output JSON only')
    parser.add_argument('--source',type=Path,required=True);parser.add_argument('--cache',type=Path,required=True)
    parser.add_argument('--format',choices=['csv','xlsx'],required=True);parser.add_argument('--max-columns',type=int,default=200);parser.add_argument('--max-sheets',type=int,default=10)
    args=parser.parse_args()
    targets=[None] if args.format=='csv' else sheet_names(args.source)[:args.max_sheets]
    profiles=[]
    for sheet in targets:
        try:profiles.append(profile_prepared(prepare_source(args.source,args.cache,ImportOptions(format=args.format,sheet=sheet)),max_columns=args.max_columns))
        except Exception as error:profiles.append({'schemaVersion':1,'sheet':sheet,'error':str(error)[:300]})
    print(json.dumps({'schemaVersion':1,'profiles':profiles}))

if __name__=='__main__':main()
