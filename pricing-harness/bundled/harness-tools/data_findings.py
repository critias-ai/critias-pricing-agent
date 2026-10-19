"""Generic data-anomaly detectors over CSV and XLSX sources. Facts only: code flags rows, the user picks treatments."""
import argparse
import csv
import hashlib
import io
import itertools
import json
import math
import re
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import duckdb
from data_workspace import ImportOptions, cached_source_hash, header_row, identifier as q, prepare_source, quote, text_encoding

VERSION = '11'
DOMINANT, RARE, KEYLIKE, RATIO, MAX_SHAPES, UNSTABLE, MIN_SHAPED = 0.95, 0.001, 0.99, 5, 30, 0.02, 5
INT, DEC = r'^[+-]?[0-9]+$', r'^[+-]?([0-9]+([.][0-9]*)?|[.][0-9]+)([eE][+-]?[0-9]+)?$'
INVISIBLE = r'[\x{200B}\x{00A0}\x{FEFF}]'
MONTH = r'^[0-9]{4}-[0-9]{2}$'  # 2025-04 is a month, read as its first day
DATA_CELL = r'^[+-]?[0-9]*[.][0-9]+$|^[0-9]{4}-[0-9]{2}-[0-9]{2}[ T][0-9]{2}:[0-9]{2}'  # a decimal or a timestamp is never a column name
stamp = lambda v: f"TRY_CAST(CASE WHEN regexp_matches({v}, '{MONTH}') THEN {v} || '-01' ELSE {v} END AS TIMESTAMP)"
LOG = 'log10(greatest(r, 1e-12))'
RATIO_LABEL = f"""CASE WHEN r < 0 THEN 'negative'
    WHEN abs(r / 12 - 1) < 0.1 THEN 'x12'
    WHEN abs(r / 24 - 1) < 0.1 THEN 'x24'
    WHEN abs(r * 12 - 1) < 0.1 THEN 'x1/12'
    WHEN abs(r * 24 - 1) < 0.1 THEN 'x1/24'
    WHEN round({LOG}) <> 0 AND abs({LOG} - round({LOG})) < 0.15 THEN 'x1e' || CAST(round({LOG}) AS INT) ELSE 'outlier' END"""
# D1 issue labels for values that fail the column's dominant parse, with a lossless rewrite where one exists.
FORMATS = [
    ('decimal-comma', r'^[+-]?[0-9]+,[0-9]{1,2}$', "replace({c}, ',', '.')"),
    ('thousands-separator', r'^[+-]?[1-9][0-9]{0,2}(,[0-9]{3})+([.][0-9]+)?$', "replace({c}, ',', '')"),
    ('currency-symbol', r'^(R\$|[£$€]) ?', "regexp_replace({c}, '^(R\\$|[£$€]) ?', '')"),
    ('parenthesised-negative', r'^\([0-9][0-9.]*\)$', "'-' || trim({c}, '()')"),
    ('unit-text', r'^[0-9.]+ *[A-Za-z]+$', "regexp_extract({c}, '^[0-9.]+')"),
    ('slash-date', r'^[0-9]{1,2}/[0-9]{1,2}/[0-9]{4}', None),
    ('serial-date', r'^[1-9][0-9]{4}([.][0-9]+)?$', "(TIMESTAMP '1899-12-30' + to_seconds(CAST(round(CAST({c} AS DOUBLE) * 86400) AS BIGINT)))::VARCHAR"),
    ('invalid-date', r'^[0-9]{4}-[0-9]{2}-[0-9]{2}', None),
    ('mixed-text', r'[0-9]', None),
]


DETECTORS = {'parse': 'ran', 'shape': 'skipped: no text columns', 'domain': 'skipped: no key shared with another source',
             'uniqueness': 'ran', 'ratio': 'skipped: no key whose groups describe its numbers', 'identity': 'skipped: no decimal column with two other numbers, or too many numeric columns',
             'time': 'skipped: no date columns'}


def unreadable(e):
    if isinstance(e, UnicodeDecodeError):
        return 'not UTF-8 or UTF-16; convert it to UTF-8 first'
    if isinstance(e, csv.Error):
        return 'delimiter not recognised (only , ; tab and | are read); convert it first'
    return f'{type(e).__name__}: {e}'[:200]


def prepare(path, cache):
    """Each CSV, or each sheet of a workbook, as its own source; a source that cannot be read becomes its reason."""
    try:
        sheets, row = [None], 1
        if not path.lower().endswith(('.xlsx', '.xlsm')):  # a separator such as +| sniffs as | and leaves + on every value
            with open(path, encoding=text_encoding(path)) as f:
                sample = f.read(65536)
            if any(sample.count(d) and sum(sample.count(x + d) for x in '+&~^#*!') == sample.count(d) for d in ',;\t|'):  # +| fields, &| row ends
                raise csv.Error('multi-character delimiter')
            delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
            row = header_row(sample, delimiter)  # a title or date-range row above the header is left out
            first = next(itertools.islice(csv.reader(io.StringIO(sample), delimiter=delimiter), row - 1, None), [])
            if any(re.match(DATA_CELL, v.strip()) for v in first):
                return {path: f"no header row: the first row holds data ({', '.join(first[:3])}, …); add column names first"[:200]}
        else:
            from openpyxl import load_workbook
            sheets = load_workbook(path, read_only=True).sheetnames
        return {path if s is None else f'{path}#{s}': prepare_source(Path(path), cache, ImportOptions(format='csv' if s is None else 'xlsx', sheet=s, encoding='utf-8-sig', header_row=row if s is None else 1))  # utf-8-sig: a byte-order mark is not part of the first header
                for s in sheets}
    except FileNotFoundError:
        raise
    except Exception as e:
        return {path: unreadable(e)}


def detect(paths, cache):
    cache = Path(cache)
    key = hashlib.sha256(json.dumps([VERSION, *sorted((str(p), cached_source_hash(Path(p), cache)) for p in paths)]).encode()).hexdigest()[:16]
    out = cache / f'findings-{key}'
    if (out / 'findings.json').is_file():
        return out, json.loads((out / 'findings.json').read_text())
    prepared = {name: p for path in paths for name, p in prepare(str(path), cache).items()}
    db = duckdb.connect()
    db.execute('CREATE TABLE flags (finding VARCHAR, rid BIGINT)')
    src, skipped = {}, {name: p for name, p in prepared.items() if isinstance(p, str)}
    for i, (path, p) in enumerate((n, p) for n, p in prepared.items() if n not in skipped):
        db.execute(f'ATTACH {quote(p.database_path)} AS s{i} (READ_ONLY)')
        src[path] = {'t': f's{i}.raw', 'rows': p.row_count, 'cols': [r[0] for r in db.execute(f'DESCRIBE s{i}.raw').fetchall()]}
    findings, ran = [], {path: {'parse', 'uniqueness'} for path in src}

    def add(path, column, detector, issue, rowids, fix=None, exact=True):
        s, fid = src[path], 'F-' + hashlib.sha1(json.dumps([detector, path, column, issue]).encode()).hexdigest()
        db.execute(f'INSERT INTO flags SELECT {quote(fid)}, rid FROM ({rowids})')
        flagged = f"FROM {s['t']} WHERE rowid IN (SELECT rid FROM flags WHERE finding = {quote(fid)})"
        rows, = db.execute(f'SELECT count(*) {flagged}').fetchone()
        if not rows:
            return
        value = db.execute(f"SELECT sum(abs(TRY_CAST({q(s['value'])} AS DOUBLE))) FILTER (WHERE isfinite(TRY_CAST({q(s['value'])} AS DOUBLE))) {flagged}").fetchone()[0] if s['value'] else None
        examples = [r[0] for r in db.execute(f"SELECT DISTINCT {q(s['cols'][0] if column == '*' else column)} {flagged} ORDER BY 1 LIMIT 3").fetchall()]
        findings.append({'id': fid, 'detector': detector, 'source': path, 'column': column, 'issue': issue, 'rows': rows,
                         'valueAtStake': None if value is None or math.isinf(value) else round(value, 2), 'valueColumn': s['value'],
                         'examples': examples, 'fix': fix and fix.replace('{c}', q(column)), 'exact': exact})

    # Column facts shared by every detector: parse shares, uniqueness, code-likeness.
    for path, s in src.items():
        s['facts'] = {}
        for c in s['cols']:
            v = f"NULLIF(trim({q(c)}), '')"
            present, ints, decs, stamps, distinct, codey, total = db.execute(f"""SELECT count({v}),
                count(*) FILTER (WHERE regexp_matches({v}, {quote(INT)})), count(*) FILTER (WHERE regexp_matches({v}, {quote(DEC)})),
                count({stamp(v)}), count(DISTINCT {v}),
                count(*) FILTER (WHERE length({v}) <= 20 AND NOT regexp_matches({v}, ' ')), sum(abs(TRY_CAST({v} AS DOUBLE))) FILTER (WHERE isfinite(TRY_CAST({v} AS DOUBLE)))
                FROM {s['t']}""").fetchone()
            s['facts'][c] = dict(present=present, ints=ints, decs=decs, stamps=stamps, distinct=distinct, total=total or 0, codey=codey,
                                 numeric=bool(present) and decs >= DOMINANT * present,
                                 stamp=bool(present) and decs < DOMINANT * present and stamps >= DOMINANT * present,
                                 unique=present >= 0.5 * s['rows'] and distinct >= KEYLIKE * present,
                                 code=bool(present) and decs < DOMINANT * present and stamps < DOMINANT * present and codey >= 0.99 * present and distinct >= 10)
        measures = [c for c, f in s['facts'].items() if f['numeric'] and not f['unique']]
        decimals = [c for c in measures if 2 * s['facts'][c]['ints'] < s['facts'][c]['decs']]
        # shortcut: whole-number measures (unit counts) stand in only when a source has no decimal column; may pick an ID column.
        s['values'] = sorted(decimals or measures, key=lambda c: -s['facts'][c]['total'])
        s['value'] = s['values'][0] if s['values'] else None

    # D1 parse: values failing the column's dominant parse, labelled by the format they do match. A column wholly in one
    # non-standard format (currency text, decimal commas, slash dates) has no standard parse, so all of it is flagged.
    for path, s in src.items():
        for c, f in s['facts'].items():
            v, t = f"NULLIF(trim({q(c)}), '')", s['t']
            label = ' '.join(f"WHEN regexp_matches({v}, {quote(rx)}) THEN {quote(name)}" for name, rx, _ in FORMATS)
            label = f"CASE WHEN {v} = {quote(c)} THEN 'header-text' {label} ELSE 'text' END"
            if f['numeric'] or f['stamp']:
                ok = f"regexp_matches({v}, {quote(DEC)})" if f['numeric'] else f'{stamp(v)} IS NOT NULL'
                bad, kind = f'FROM {t} WHERE {v} IS NOT NULL AND NOT {ok}', 'DOUBLE' if f['numeric'] else 'TIMESTAMP'
            else:
                top = db.execute(f'SELECT {label} l, count(*) n FROM {t} WHERE {v} IS NOT NULL GROUP BY 1 ORDER BY 2 DESC, 1 LIMIT 1').fetchone()
                if not top or top[0] in ('text', 'mixed-text', 'header-text') or top[1] < DOMINANT * f['present']:
                    continue
                bad, kind = f'FROM {t} WHERE {v} IS NOT NULL', 'TIMESTAMP' if 'date' in top[0] else 'DOUBLE'
            for issue, in db.execute(f'SELECT DISTINCT {label} {bad}').fetchall():
                fix = next((fx for name, _, fx in FORMATS if name == issue), None)
                pick = f'{bad} AND {label} = {quote(issue)}'
                if fix and db.execute(f"SELECT count(*) {bad} AND CASE WHEN {label} = {quote(issue)} THEN TRY_CAST({fix.replace('{c}', q(c))} AS {kind}) IS NULL END").fetchone()[0]:
                    fix = None
                if issue == 'slash-date':  # which part is the day decides the date; say when the data cannot tell
                    first, second = db.execute(f"SELECT max(TRY_CAST(split_part({v}, '/', 1) AS INT)), max(TRY_CAST(split_part({v}, '/', 2) AS INT)) {pick}").fetchone()
                    issue += ' ' + ('mixed-order' if first > 12 and second > 12 else 'day-first' if first > 12 else 'month-first' if second > 12 else 'ambiguous')
                add(path, c, 'parse', issue, f'SELECT rowid AS rid {pick}', fix)
            if f['numeric'] and f['decs'] > f['ints'] >= 0.999 * f['decs']:  # an integer column holding a few decimals
                add(path, c, 'parse', 'float-formatted', f"SELECT rowid AS rid FROM {t} WHERE regexp_matches({v}, '^[+-]?[0-9]+[.]0+$')",
                    r"regexp_replace({c}, '[.]0+$', '')")
                add(path, c, 'parse', 'non-integer', f"SELECT rowid AS rid FROM {t} WHERE regexp_matches({v}, {quote(DEC)}) AND NOT regexp_matches({v}, {quote(INT)}) AND NOT regexp_matches({v}, '^[+-]?[0-9]+[.]0+$')")

    # D2 shape: flag spelling variants of any text value, and in code-like columns of MIN_SHAPED+ values, shapes that are
    # rare or outside the common shapes that hold 95% of values; normalization could merge distinct records, so offer no fix.
    for path, s in src.items():
        for c, f in s['facts'].items():
            if f['numeric'] or f['stamp'] or not f['present']:
                continue
            ran[path].add('shape')
            t, norm = s['t'], f"upper(trim(regexp_replace({q(c)}, {quote(INVISIBLE)}, '', 'g')))"
            # Spellings that collide after normalising: the most frequent is canonical (ties by spelling), the rest are variants.
            add(path, c, 'shape', 'variant-of-common-value',
                f'''SELECT rowid AS rid FROM {t} WHERE {q(c)} IN (SELECT v FROM (
                    SELECT v, row_number() OVER (PARTITION BY n ORDER BY k DESC, v) rn, count(*) OVER (PARTITION BY n) spellings
                    FROM (SELECT {q(c)} v, {norm} n, count(*) k FROM {t} WHERE {q(c)} IS NOT NULL GROUP BY 1, 2)) WHERE spellings > 1 AND rn > 1)''')
            if f['distinct'] < MIN_SHAPED or f['codey'] < DOMINANT * f['present']:  # plan names, countries, free text: shapes mean nothing
                continue
            shape = f"regexp_replace(regexp_replace({q(c)}, '[0-9]', '9', 'g'), '[A-Za-z]', 'A', 'g')"
            shapes = db.execute(f"SELECT {shape} sh, count(*) n FROM {t} WHERE {q(c)} IS NOT NULL GROUP BY sh ORDER BY n DESC, sh").fetchall()
            common = {sh for (sh, _), held in zip(shapes, itertools.accumulate([0] + [n for _, n in shapes])) if held < DOMINANT * f['present']}
            for sh, n in shapes if len(shapes) <= MAX_SHAPES else []:  # many shapes means free-form values, not codes
                if n >= RARE * f['present'] and sh in common:  # common: the most frequent shapes that together hold 95% of values
                    continue
                add(path, c, 'shape', f'rare-shape {sh}', f'SELECT rowid AS rid FROM {t} WHERE {shape} = {quote(sh)}')

    # D3 domain: values of a column missing from the source where they are a key, and where they do occur. A column
    # refers to a key by sharing its name, or, under another name, when most of its distinct values are that key's.
    masters = {}
    for c in {c for s in src.values() for c in s['cols']}:
        keyed = [p for p, s in src.items() if c in s['facts'] and s['facts'][c]['unique']]
        if keyed and sum(c in s['cols'] for s in src.values()) > 1:
            masters[c] = max(keyed, key=lambda p: (src[p]['rows'], p))
    pairs = [(path, c, master, c) for c, master in masters.items() for path, s in src.items() if path != master and c in s['cols']]
    # shortcut: only code-like text columns are matched by value; numeric IDs under different names still need a shared name.
    keyset = [(p, k) for p, s in src.items() for k, f in s['facts'].items() if f['unique'] and f['code']]
    for path, s in src.items():
        for c in (c for c, f in s['facts'].items() if f['code'] and not f['unique'] and c not in masters):
            sample = f"(SELECT DISTINCT {q(c)} v FROM {s['t']} WHERE {q(c)} IS NOT NULL LIMIT 1000)"
            for mp, k in keyset:
                if mp != path and k != c and db.execute(f"SELECT count(*) >= 0.5 * (SELECT count(*) FROM {sample}) FROM {sample} WHERE v IN (SELECT {q(k)} FROM {src[mp]['t']})").fetchone()[0]:
                    pairs.append((path, c, mp, k))
    lookups = [(p, c) for p in {m for _, _, m, _ in pairs} for c, f in src[p]['facts'].items() if not (f['numeric'] or f['stamp'])]
    for path, c, master, k in pairs:
        s, keys = src[path], f"(SELECT {q(k)} FROM {src[master]['t']} WHERE {q(k)} IS NOT NULL)"
        ran[path].add('domain'), ran[master].add('domain')
        missing = f"FROM {s['t']} WHERE NULLIF(trim({q(c)}), '') IS NOT NULL AND {q(c)} NOT IN {keys}"
        if not 0 < db.execute(f'SELECT count(DISTINCT {q(c)}) {missing}').fetchone()[0] <= s['facts'][c]['distinct'] / 2:
            continue
        values = f'(SELECT DISTINCT {q(c)} {missing})'
        best = max(((db.execute(f"SELECT count(DISTINCT {q(oc)}) FROM {src[op]['t']} WHERE {q(oc)} IN {values}").fetchone()[0], op, oc)
                    for op, oc in lookups if oc != c), default=(0, None, None))
        if best[0]:
            found = f"(SELECT {q(best[2])} FROM {src[best[1]]['t']} WHERE {q(best[2])} IS NOT NULL)"
            add(path, c, 'domain', f'value-of {best[1]}:{best[2]}', f'SELECT rowid AS rid {missing} AND {q(c)} IN {found}')
            missing += f' AND {q(c)} NOT IN {found}'
        add(path, c, 'domain', f'missing-from {master}' + ('' if k == c else f':{k}'), f'SELECT rowid AS rid {missing}')

    # D4 uniqueness: repeated rows, conflicting duplicate keys, and whole documents repeated under new keys.
    for path, s in src.items():
        t, every = s['t'], ', '.join(q(c) for c in s['cols'])
        add(path, '*', 'uniqueness', 'repeated-row', f'SELECT rid FROM (SELECT rowid AS rid, row_number() OVER (PARTITION BY {every} ORDER BY rowid) n FROM {t}) WHERE n > 1')
        for c in (c for c, f in s['facts'].items() if f['unique'] and f['distinct'] < f['present']):
            add(path, c, 'uniqueness', 'duplicate-key conflicting',
                f'SELECT rowid AS rid FROM {t} WHERE {q(c)} IN (SELECT {q(c)} FROM {t} GROUP BY 1 HAVING count(DISTINCT hash({every})) > 1)')
        docs = [c for c, f in s['facts'].items() if not f['unique'] and f['distinct'] >= 0.01 * f['present'] > 0 and (f['code'] or f['ints'] >= DOMINANT * f['present'])]
        for d in docs:
            rest = ', '.join(q(c) for c, f in s['facts'].items() if not f['unique'] and c != d)
            if not rest: continue
            add(path, d, 'uniqueness', 'repeated-document', f"""SELECT rowid AS rid FROM {t} WHERE {q(d)} IN (
                SELECT doc FROM (SELECT doc, fp, min(doc) OVER (PARTITION BY fp) earliest FROM (
                    SELECT {q(d)} doc, string_agg(hash({rest})::VARCHAR, ',' ORDER BY hash({rest})) fp FROM {t} GROUP BY 1 HAVING count(*) > 1)) WHERE doc <> earliest)""",
                exact=False)

    # D7 peer ratio: values a power of ten, twelve or a sign away from their key group's median, in columns those
    # medians describe well (the best-fitting key leaves under UNSTABLE of rows far from the median, not counting rows a
    # dozen or two dozen out, so a common each-vs-case error cannot switch the check off). Those column and
    # key pairs are also kept as `peers`, so the SQL session can flag such rows in what a statement sums.
    peers = []
    for path, s in src.items():
        groups = [c for c in s['cols'] if (c in masters and masters[c] != path or s['facts'][c]['code']) and 3 * s['facts'][c]['distinct'] <= s['facts'][c]['present']]
        numbers = [c for c, f in s['facts'].items() if f['numeric'] and not f['unique'] and c not in masters]
        for m in numbers:
            fits = []
            for g in groups:
                ratio = f"""SELECT rowid AS rid, val / med AS r FROM (SELECT rowid, {q(g)} g, TRY_CAST({q(m)} AS DOUBLE) val FROM {s['t']}) x
                    JOIN (SELECT {q(g)} g, median(TRY_CAST({q(m)} AS DOUBLE)) med FROM {s['t']}
                        WHERE isfinite(TRY_CAST({q(m)} AS DOUBLE)) GROUP BY 1 HAVING count({q(m)}) >= 3 AND med > 0) USING (g)
                    WHERE val <> 0 AND isfinite(val) AND (val / med < {1 / RATIO} OR val / med > {RATIO})"""
                fits.append((db.execute(f"SELECT count(*) FROM ({ratio}) WHERE {RATIO_LABEL} NOT IN ('x12', 'x24', 'x1/12', 'x1/24')").fetchone()[0], g not in masters, s['facts'][g]['distinct'], g, ratio))  # ties: another source's key, then fewer, larger groups
            far, _, _, g, ratio = min(fits, default=(math.inf, None, None, None, None))
            if far >= UNSTABLE * s['facts'][m]['present']:
                continue
            peers.append({'source': path, 'column': m, 'key': g})
            ran[path].add('ratio')
            for label, in db.execute(f"SELECT DISTINCT {RATIO_LABEL} FROM ({ratio})").fetchall():
                add(path, m, 'ratio', f'{label} vs {g} median', f'SELECT rid FROM ({ratio}) WHERE {RATIO_LABEL} = {quote(label)}', exact=False)

    # D5 identity: a decimal column equal to the product of two others on nearly every informative row (value not zero,
    # neither factor 0 or ±1, where any product would hold trivially); flag every row where it is not.
    for path, s in src.items():
        nums = [c for c, f in s['facts'].items() if f['numeric'] and not f['unique']]
        triples = [(a, b, c) for c in nums if 2 * s['facts'][c]['ints'] < s['facts'][c]['decs'] for a, b in itertools.combinations([n for n in nums if n != c], 2)]
        if not 0 < len(triples) <= 3000:  # shortcut: very wide numeric sources are skipped (reported in coverage); batch them if one matters
            continue
        x = lambda c: f'TRY_CAST({q(c)} AS DOUBLE)'
        off = lambda a, b, c: f'abs({x(a)} * {x(b)} - {x(c)}) > 0.01 + 0.001 * abs({x(c)})'
        told = lambda a, b, c: f'{x(c)} <> 0 AND abs({x(a)}) NOT IN (0, 1) AND abs({x(b)}) NOT IN (0, 1)'
        counts = [n for i in range(0, len(triples), 200) for n in db.execute('SELECT ' + ', '.join(  # batches bound memory on wide sources
            f'count(*) FILTER (WHERE {told(a, b, c)}), count(*) FILTER (WHERE {told(a, b, c)} AND {off(a, b, c)}), count(*) FILTER (WHERE {off(a, b, c)})'
            for a, b, c in triples[i:i + 200]) + f" FROM {s['t']}").fetchone()]
        ran[path].add('identity')
        for (a, b, c), both, wrong, failing in zip(triples, counts[::3], counts[1::3], counts[2::3]):
            if both >= 0.2 * s['rows'] and wrong <= (1 - DOMINANT) * both and failing:
                add(path, c, 'identity', f'not {a} × {b}', f"SELECT rowid AS rid FROM {s['t']} WHERE {off(a, b, c)}")

    # D6 time: a few dates after today, and among rows where two dates differ, the few whose order is the reverse of the
    # rest; the second date may sit in this row or, through a key, in the row it refers to (signed up, churned).
    def order(path, frm, pairs):  # (column, issue, ts, other ts): flag the few rows where ts > other ts
        for c, issue, ta, tb in pairs:
            both, ahead = db.execute(f"SELECT count(*) FILTER (WHERE {ta} <> {tb}), count(*) FILTER (WHERE {ta} > {tb}) {frm}").fetchone()
            if both and 0 < ahead <= (1 - DOMINANT) * both:
                add(path, c, 'time', issue, f"SELECT x.rowid AS rid {frm} AND {ta} > {tb}")

    stamps = lambda p: [c for c, f in src[p]['facts'].items() if f['stamp']]
    x, y = lambda c: stamp('x.' + q(c)), lambda c: stamp('y.' + q(c))
    day = lambda e: f'CAST({e} AS DATE)'  # across tables a timestamp on the last day is not after that date
    for path, s in src.items():
        ran[path].update(['time'] if stamps(path) else [])
        for c in stamps(path):
            if db.execute(f"SELECT count(*) FILTER (WHERE {stamp(q(c))} > current_date + 1) <= {1 - DOMINANT} * count({stamp(q(c))}) FROM {s['t']}").fetchone()[0]:
                add(path, c, 'time', 'after today', f"SELECT rowid AS rid FROM {s['t']} WHERE {stamp(q(c))} > current_date + 1")
        order(path, f"FROM {s['t']} x WHERE true", [(a, f'after {b}', x(a), x(b)) for a, b in itertools.permutations(stamps(path), 2)])
    for path, c, master, k in (p for p in pairs if not src[p[2]]['facts'][p[3]]['numeric']):  # a count that happens to be unique is no key
        order(path, f"FROM {src[path]['t']} x JOIN {src[master]['t']} y ON x.{q(c)} = y.{q(k)} WHERE true",
              [p for a in stamps(path) for b in stamps(master) for p in ((a, f'after {master}:{b}', day(x(a)), day(y(b))), (a, f'before {master}:{b}', day(y(b)), day(x(a))))])
    # shortcut: overlapping validity intervals are not checked; that needs to know which column the intervals belong to.

    coverage = [{'source': p, 'status': 'ran', 'valueColumns': s['values'], 'detectors': {d: 'ran' if d in ran[p] else why for d, why in DETECTORS.items()}}
                for p, s in src.items()] + [{'source': p, 'status': 'skipped', 'reason': why, 'valueColumns': [], 'detectors': {d: 'skipped: not read' for d in DETECTORS}}
                                            for p, why in skipped.items()]
    out.mkdir(parents=True, exist_ok=True)
    db.execute(f"COPY flags TO {quote(out / 'flags.parquet')} (FORMAT parquet)")
    result = {'schemaVersion': 1, 'version': VERSION, 'ratio': RATIO, 'peers': peers, 'coverage': coverage, 'findings': sorted(findings, key=lambda f: (-(f['valueAtStake'] or 0), f['id']))}
    (out / 'findings.json').write_text(json.dumps(result, indent=1))
    return out, result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Detect data anomalies across CSV and XLSX sources; output JSON only')
    parser.add_argument('--cache', type=Path, required=True)
    parser.add_argument('--source', action='append', required=True)
    args = parser.parse_args()
    folder, result = detect(args.source, args.cache)
    print(json.dumps({'findingsDir': str(folder), **result}))
